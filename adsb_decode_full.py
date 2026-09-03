"""Decodificar TODO lo que trae cada mensaje, y volcarlo a un Excel por columnas.

La diferencia con adsb_record.py es el criterio, no la tecnica. Ese grabador
extrae los 8 campos que el proyecto necesita para contar operaciones y tira el
resto; sirve para operar. Esto no descarta nada: toma el diccionario completo
que devuelve pyModeS y lo vuelca entero, campo por campo. Sirve para MIRAR --
para decidir, viendo los datos, que analisis vale la pena construir.

Las columnas NO estan escritas a mano. Se juntan de lo que efectivamente
aparecio en el aire (la union de las claves de todos los mensajes decodificados)
y despues se ordenan segun adsb_catalogo.ORDEN. Es a proposito: una lista fija
de columnas garantiza que el dia que pyModeS agregue un campo, o que aparezca
una aeronave que emita algo que aca no se previo, el dato se pierda en silencio
-- que es exactamente el problema que este modulo existe para terminar. Si
aparece un campo que el catalogo no describe, la columna igual se exporta y el
diccionario la marca como "sin catalogar".

Entra el hex crudo de adsb_raw.py y no la antena en vivo, tambien a proposito:
el archivo crudo se puede volver a decodificar cuantas veces se quiera, con el
decodificador de hoy o con el de dentro de seis meses. Un pipeline que decodifica
en vivo solo produce lo que sabia extraer en el momento.

Uso:
  python adsb_decode_full.py                        el crudo de hoy -> Excel
  python adsb_decode_full.py --raw archivo.csv
  python adsb_decode_full.py --out mi_analisis.xlsx
"""
from __future__ import annotations

import argparse
import csv
import os
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import pyModeS as pms

import aircraft_db
from adsb_catalogo import BDS_NOMBRE, DF_NOMBRE, ORDEN, describir
from adsb_raw import RAW_DIR
from receiver import RECEIVER_NAME, distance_km, horizonte_km, surface_ref_default

SALIDA_DIR = Path(__file__).parent / "output" / "adsb"


def _limpiar(valor):
    """Un valor que Excel pueda guardar. Las listas y tuplas se hacen texto."""
    if isinstance(valor, (list, tuple, set)):
        return ", ".join(str(v) for v in valor)
    if isinstance(valor, (int, float, str, bool)) or valor is None:
        return valor
    return str(valor)


def decodificar(raw_path: Path, *, con_lookup: bool = True) -> tuple[list[dict], list[str]]:
    """Cada mensaje del archivo crudo, con todos sus campos. (filas, columnas).

    El decodificador es el MISMO PipeDecoder que usa la grabacion en vivo, con
    la misma surface_ref: si diera algo distinto, este Excel no describiria el
    sistema sino a si mismo.
    """
    decoder = pms.PipeDecoder(surface_ref=surface_ref_default())
    filas: list[dict] = []
    vistas: set[str] = set()
    # Cache de matriculas: una consulta por aeronave y no por mensaje. En un
    # archivo de 20 000 mensajes de 200 aviones eso son 200 consultas y no
    # 20 000, y la diferencia se nota a simple vista.
    cache: dict[str, dict] = {}

    with raw_path.open(encoding="utf-8", newline="") as fh:
        for cruda in csv.DictReader(fh):
            try:
                epoch = float(cruda["epoch"])
            except (KeyError, TypeError, ValueError):
                continue
            hex_msg = (cruda.get("hex") or "").strip()
            if not hex_msg:
                continue

            try:
                decodificado = dict(decoder.decode(hex_msg, timestamp=epoch))
            except Exception as exc:
                # Un mensaje que revienta al decodificar NO se descarta: se
                # anota. Descartarlo dejaria el archivo con menos filas que
                # mensajes y nadie sabria por que. La columna existe para que
                # un problema sistematico del decodificador sea visible en el
                # Excel en vez de deducirse de un faltante.
                decodificado = {"error_decodificacion": f"{type(exc).__name__}: {exc}"}

            fila = {
                "utc": datetime.fromtimestamp(epoch, timezone.utc)
                               .isoformat(timespec="milliseconds"),
                "epoch": round(epoch, 3),
                "hex": hex_msg,
            }
            # Vacio y no 0.0: el camino de rtl_adsb no mide la senal, y un 0
            # ahi seria "llego a fondo de escala", que es lo contrario.
            crudo_dbfs = (cruda.get("dbfs") or "").strip()
            if crudo_dbfs:
                try:
                    fila["dbfs"] = float(crudo_dbfs)
                except ValueError:
                    pass
            fila.update({k: _limpiar(v) for k, v in decodificado.items()})

            # --- derivados ---
            df = decodificado.get("df")
            if df is not None:
                fila["df_nombre"] = DF_NOMBRE.get(df, f"DF{df}")
            bds = decodificado.get("bds")
            if bds:
                fila["bds_nombre"] = BDS_NOMBRE.get(bds, "")

            icao = (decodificado.get("icao") or "").lower()
            if icao and con_lookup:
                if icao not in cache:
                    cache[icao] = aircraft_db.lookup(icao) or {}
                entrada = cache[icao]
                if entrada:
                    fila["registration"] = entrada.get("registration")
                    fila["aircraft_type"] = aircraft_db.describe_type(entrada)
                    fila["operator"] = entrada.get("operator")

            lat, lon = decodificado.get("latitude"), decodificado.get("longitude")
            if lat is not None and lon is not None:
                km = distance_km(lat, lon)
                fila["distancia_km"] = round(km, 2) if km is not None else None
                # El horizonte depende de la altitud: un avion a 35 000 ft se
                # ve a 400 km y uno en pista a 10. Comparar contra un radio
                # fijo daria falsos positivos arriba y dejaria pasar lo de
                # abajo, que es justo donde estan los errores de CPR.
                alt = decodificado.get("altitude")
                if km is not None and alt is not None:
                    fila["sobre_horizonte"] = km > horizonte_km(float(alt))

            vistas.update(fila.keys())
            filas.append(fila)

    # ORDEN primero, y al final lo que aparecio pero el catalogo no previo.
    # Ese resto NO se descarta ni se esconde: si la antena lo trajo, va.
    columnas = [c for c in ORDEN if c in vistas]
    columnas += sorted(vistas - set(columnas))
    return filas, columnas


def resumen_por_aeronave(filas: list[dict]) -> list[dict]:
    """Una fila por aeronave: que se supo de cada una y de cuantos mensajes."""
    porico: dict[str, list[dict]] = defaultdict(list)
    for fila in filas:
        if fila.get("icao"):
            porico[str(fila["icao"]).lower()].append(fila)

    def _numeros(msgs, campo):
        out = []
        for m in msgs:
            v = m.get(campo)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                out.append(float(v))
        return out

    def _primero(msgs, campo):
        for m in msgs:
            if m.get(campo) not in (None, ""):
                return m[campo]
        return None

    resumen = []
    for icao, msgs in sorted(porico.items(), key=lambda kv: -len(kv[1])):
        epochs = _numeros(msgs, "epoch")
        alts = _numeros(msgs, "altitude")
        dists = _numeros(msgs, "distancia_km")
        familias = Counter(m.get("bds_nombre") or m.get("df_nombre") or "?"
                           for m in msgs)
        resumen.append({
            "icao": icao,
            "registration": _primero(msgs, "registration"),
            "aircraft_type": _primero(msgs, "aircraft_type"),
            "operator": _primero(msgs, "operator"),
            "callsign": _primero(msgs, "callsign"),
            "mensajes": len(msgs),
            "primero_utc": min((m["utc"] for m in msgs), default=None),
            "ultimo_utc": max((m["utc"] for m in msgs), default=None),
            "segundos_oida": round(max(epochs) - min(epochs), 1) if epochs else None,
            "con_posicion": sum(1 for m in msgs if m.get("latitude") is not None),
            "altitud_min_ft": min(alts) if alts else None,
            "altitud_max_ft": max(alts) if alts else None,
            "distancia_max_km": round(max(dists), 2) if dists else None,
            "squawk": _primero(msgs, "squawk"),
            "en_tierra_declarado": any(
                m.get("bds") == "0,6" or m.get("vertical_status") == "on-ground"
                for m in msgs),
            "familias": "; ".join(f"{k} x{v}" for k, v in familias.most_common()),
        })
    return resumen


def exportar_excel(filas: list[dict], columnas: list[str], destino: Path,
                   raw_path: Path, nota: str | None = None) -> Path:
    """Cuatro hojas: los mensajes, que significa cada columna, por aeronave, y
    de donde salio todo.

    El diccionario va DENTRO del archivo y no en un README aparte porque el
    archivo se va a abrir solo, seis meses despues, en otra maquina, y ahi la
    pregunta va a ser "que es nuc_p". Un Excel que no se explica a si mismo
    obliga a adivinar, y adivinar sobre datos es como salen los informes
    equivocados.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    encabezado = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="1F3864")

    # --- hoja 1: los mensajes ---
    ws = wb.active
    ws.title = "mensajes"
    ws.append(columnas)
    for celda in ws[1]:
        celda.font, celda.fill = encabezado, fondo
        celda.alignment = Alignment(horizontal="center")
    for fila in filas:
        # Los valores van con su TIPO, no como texto: asi Excel los trata como
        # numeros y no repite el desastre de leer -34.66 como -3.466 por el
        # separador decimal en castellano. Ese problema es de importar CSV;
        # el xlsx guarda el numero binario y no lo tiene.
        ws.append([fila.get(c) for c in columnas])
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(col) + 3, 11), 26)

    # --- hoja 2: el diccionario ---
    wd = wb.create_sheet("diccionario")
    wd.append(["columna", "que es", "unidad", "de que mensaje sale",
               "filas con dato", "% de filas"])
    for celda in wd[1]:
        celda.font, celda.fill = encabezado, fondo
    total = len(filas) or 1
    for col in columnas:
        que, unidad, origen = describir(col)
        con_dato = sum(1 for f in filas if f.get(col) not in (None, ""))
        wd.append([col, que, unidad, origen, con_dato,
                   round(100.0 * con_dato / total, 1)])
    wd.freeze_panes = "A2"
    for letra, ancho in (("A", 24), ("B", 62), ("C", 12), ("D", 30),
                         ("E", 14), ("F", 11)):
        wd.column_dimensions[letra].width = ancho
    for fila in wd.iter_rows(min_row=2, min_col=2, max_col=2):
        fila[0].alignment = Alignment(wrap_text=True, vertical="top")

    # --- hoja 3: por aeronave ---
    resumen = resumen_por_aeronave(filas)
    wa = wb.create_sheet("aeronaves")
    if resumen:
        cols_a = list(resumen[0].keys())
        wa.append(cols_a)
        for celda in wa[1]:
            celda.font, celda.fill = encabezado, fondo
        for r in resumen:
            wa.append([r.get(c) for c in cols_a])
        wa.freeze_panes = "B2"
        wa.auto_filter.ref = wa.dimensions
        for i, col in enumerate(cols_a, start=1):
            wa.column_dimensions[get_column_letter(i)].width = (
                50 if col == "familias" else min(max(len(col) + 3, 12), 22))

    # --- hoja 4: de donde salio esto ---
    wp = wb.create_sheet("procedencia")
    wp.append(["dato", "valor"])
    for celda in wp[1]:
        celda.font, celda.fill = encabezado, fondo
    epochs = [f["epoch"] for f in filas if isinstance(f.get("epoch"), (int, float))]
    for clave, valor in [
        # La procedencia va PRIMERA y en el archivo mismo. Un Excel con datos
        # de referencia y uno con lo que oyo la antena se ven identicos, y
        # confundirlos es peor que no tener ninguno: uno describe el mundo y
        # el otro describe el formato.
        ("QUE ES ESTO", nota or "Capturado por la antena de este proyecto."),
        ("archivo crudo", str(raw_path)),
        ("receptor", RECEIVER_NAME),
        ("ADSB_RECEIVER", os.environ.get("ADSB_RECEIVER") or
         "(NO DEFINIDA -- las distancias salen del default del codigo, que "
         "puede no ser donde esta la antena)"),
        ("mensajes decodificados", len(filas)),
        ("aeronaves distintas", len(resumen)),
        ("columnas exportadas", len(columnas)),
        ("desde (UTC)", min((f["utc"] for f in filas), default="-")),
        ("hasta (UTC)", max((f["utc"] for f in filas), default="-")),
        ("duracion (s)", round(max(epochs) - min(epochs), 1) if epochs else "-"),
        ("decodificador", f"pyModeS {getattr(pms, '__version__', '?')}"),
        ("generado (UTC)", datetime.now(timezone.utc).isoformat(timespec="seconds")),
    ]:
        wp.append([clave, valor])
    wp.column_dimensions["A"].width = 26
    wp.column_dimensions["B"].width = 60

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--raw", type=Path, default=None,
                    help="archivo de hex crudo; por defecto el de hoy")
    ap.add_argument("--out", type=Path, default=None)
    ap.add_argument("--sin-lookup", action="store_true",
                    help="no consultar matriculas en la base OpenSky")
    ap.add_argument("--nota", default=None,
                    help="que es este archivo; va en la hoja de procedencia")
    args = ap.parse_args()

    raw = args.raw
    if raw is None:
        dia = datetime.now().strftime("%Y-%m-%d")
        raw = RAW_DIR / f"raw_{dia}.csv"
    if not raw.exists():
        print(f"No existe {raw}.\nGraba primero con:  python adsb_raw.py "
              f"--seconds 300", file=sys.stderr)
        raise SystemExit(1)

    filas, columnas = decodificar(raw, con_lookup=not args.sin_lookup)
    if not filas:
        # Vacio se explica, no se rellena.
        print(f"{raw} no tiene mensajes decodificables. El Excel no se genera: "
              f"un archivo con encabezados y sin filas se parece demasiado a "
              f"un analisis que dio cero.", file=sys.stderr)
        raise SystemExit(1)

    destino = args.out or (SALIDA_DIR / f"{raw.stem}_completo.xlsx")
    exportar_excel(filas, columnas, destino, raw, nota=args.nota)
    print(f"{len(filas)} mensajes, {len(columnas)} columnas -> {destino}")


if __name__ == "__main__":
    main()
